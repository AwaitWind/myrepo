#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Estimate rendered page count of the generated IEEE docx using real
Times New Roman metrics + Word's line-wrapping / single-spacing rules."""

import sys

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from PIL import ImageFont

FDIR = "/System/Library/Fonts/Supplemental/"
FACES = {
    (False, False): FDIR + "Times New Roman.ttf",
    (True, False): FDIR + "Times New Roman Bold.ttf",
    (False, True): FDIR + "Times New Roman Italic.ttf",
    (True, True): FDIR + "Times New Roman Bold Italic.ttf",
}
_cache = {}


def face(bold, italic, size):
    key = (bold, italic, round(size, 2))
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FACES[(bold, italic)],
                                         max(1, int(round(size * 16))))
    return _cache[key], 16.0


def width_of(text, bold, italic, size, sub_sup=False):
    f, scale = face(bold, italic, size)
    w = f.getlength(text) / scale
    return w * (0.65 if sub_sup else 1.0)


LINE_FACTOR = 1.15   # Word "single" spacing for Times New Roman
COL_W = 3.5 * 72     # pt
FULL_W = 7.25 * 72   # pt
PAGE_H = (11 - 0.75 - 1.0) * 72   # 666 pt usable column height


def runs_of(p):
    out = []
    for r in p.runs:
        sz = r.font.size.pt if r.font.size else 10.0
        ss = bool(r.font.subscript or r.font.superscript)
        out.append((r.text, bool(r.bold), bool(r.italic), sz, ss))
    return out


def para_height(p, avail_w):
    rs = runs_of(p)
    if not rs:
        return Pt(0).pt + 10.0 * LINE_FACTOR * 0  # empty paragraph
    pf = p.paragraph_format
    first_ind = pf.first_line_indent.pt if pf.first_line_indent else 0.0
    left_ind = pf.left_indent.pt if pf.left_indent else 0.0
    max_sz = max(r[3] for r in rs)

    # build word stream, honouring explicit line breaks
    lines_text = [[]]
    for text, b, i, sz, ss in rs:
        for j, chunk in enumerate(text.split("\n")):
            if j > 0:
                lines_text.append([])
            for tok in chunk.replace("\t", " \t ").split(" "):
                if tok:
                    lines_text[-1].append((tok, b, i, sz, ss))

    total_lines = 0
    space_w = width_of(" ", False, False, max_sz)
    for toks in lines_text:
        if not toks:
            total_lines += 1
            continue
        cur = 0.0
        n = 1
        limit = avail_w - left_ind - max(0.0, first_ind)
        for tok, b, i, sz, ss in toks:
            if tok == "\t":
                cur = avail_w * 0.5   # tab jump (equations)
                continue
            w = width_of(tok, b, i, sz, ss) + space_w
            if cur + w > limit and cur > 0:
                n += 1
                cur = w
                limit = avail_w - left_ind
            else:
                cur += w
        total_lines += n

    h = total_lines * max_sz * LINE_FACTOR
    h += (pf.space_before.pt if pf.space_before else 0.0)
    h += (pf.space_after.pt if pf.space_after else 0.0)
    return h


def main(path):
    doc = Document(path)
    body = doc.element.body

    # split at the section break (title block vs two-column body)
    sect_ps = set()
    for p in doc.paragraphs:
        if p._p.find(qn("w:pPr")) is not None and \
           p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
            sect_ps.add(p._p)

    title_h = 0.0
    col_h = 0.0
    in_title = True
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph
            p = Paragraph(child, doc)
            if in_title:
                title_h += para_height(p, FULL_W)
                if child in sect_ps:
                    in_title = False
            else:
                col_h += para_height(p, COL_W)
        elif tag == "tbl":
            from docx.table import Table
            t = Table(child, doc)
            for row in t.rows:
                rh = 0.0
                for ci, cell in enumerate(row.cells):
                    cw = [0.95, 1.25, 1.25][min(ci, 2)] * 72
                    for cp in cell.paragraphs:
                        rh = max(rh, para_height(cp, cw - 6))
                col_h += rh

    cap2 = 2 * (PAGE_H - title_h) + 2 * PAGE_H
    cap_page1 = 2 * (PAGE_H - title_h)
    print(f"title block height : {title_h:6.1f} pt "
          f"({title_h/72:.2f} in, full width)")
    print(f"two-column content : {col_h:6.1f} pt of column flow")
    print(f"capacity, 2 pages  : {cap2:6.1f} pt "
          f"(p1 {cap_page1:.0f} + p2 {2*PAGE_H:.0f})")
    over = col_h - cap2
    if over <= 0:
        fill = col_h / cap2 * 100
        print(f"\n=> FITS in 2 pages, {fill:.1f}% full "
              f"({-over:.0f} pt / {-over/ (10*LINE_FACTOR):.0f} lines spare)")
    else:
        print(f"\n=> OVERFLOWS by {over:.0f} pt "
              f"(~{over/(10*LINE_FACTOR):.0f} lines, "
              f"~{over/(10*LINE_FACTOR)*8:.0f} words to cut)")
        print(f"   estimated pages: {2 + over/(2*PAGE_H):.2f}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "FQNet_IEEE_2page.docx")
