#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate a 2-page IEEE-conference-format Word document describing the
Frequency-Query (FQ-Net) surrogate + inverse-design algorithm in this folder.

    python3 make_paper.py

Output: FQNet_IEEE_2page.docx  (Times New Roman, two-column IEEE layout)
"""

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "FQNet_IEEE_2page.docx")

FONT = "Times New Roman"


# ---------------------------------------------------------------- helpers
def set_run(run, size, bold=False, italic=False, small_caps=False,
            sub=False, sup=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.small_caps = small_caps
    run.font.subscript = sub
    run.font.superscript = sup
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    return run


def tighten(p, before=0, after=0, spacing=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    return p


def set_columns(section, num=2, space_in=0.2):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space_in * 1440)))
    cols.set(qn("w:equalWidth"), "1")


def seg_para(doc, segments, size=10.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             indent=0.0, before=0, after=0):
    """segments: list of (text, style); style in
       r=roman i=italic b=bold bi=bold-italic rs/is=subscript rp/ip=superscript
    """
    p = doc.add_paragraph()
    p.alignment = align
    tighten(p, before, after)
    p.paragraph_format.first_line_indent = Inches(indent)
    for text, stl in segments:
        r = p.add_run(text)
        set_run(r, size,
                bold=stl in ("b", "bi"),
                italic=stl in ("i", "bi", "is", "ip"),
                sub=stl in ("rs", "is"),
                sup=stl in ("rp", "ip"))
    return p


def body(doc, text, indent=0.2, size=10.0):
    return seg_para(doc, [(text, "r")], size=size, indent=indent)


def h1(doc, num, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tighten(p, before=7, after=2)
    label = f"{num}.  {title}" if num else title
    set_run(p.add_run(label), 10, small_caps=True)
    return p


def h2(doc, letter, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tighten(p, before=4, after=1)
    set_run(p.add_run(f"{letter}. {title}"), 10, italic=True)
    return p


def equation(doc, lines, number):
    p = doc.add_paragraph()
    tighten(p, before=3, after=3)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(1.62), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Inches(3.40), WD_TAB_ALIGNMENT.RIGHT)
    for li, segs in enumerate(lines):
        set_run(p.add_run("\t"), 10)
        for text, stl in segs:
            set_run(p.add_run(text), 10,
                    italic=stl in ("i", "is", "ip"),
                    sub=stl in ("rs", "is"),
                    sup=stl in ("rp", "ip"))
        if li == len(lines) - 1:
            set_run(p.add_run("\t"), 10)
            set_run(p.add_run(f"({number})"), 10)
        else:
            p.add_run().add_break()
    return p


def cell_text(cell, text, size=8.0, bold=False, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    tighten(p, before=1, after=1)
    set_run(p.add_run(text), size, bold=bold, italic=italic)


def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        if edge in kwargs:
            el = borders.find(qn(f"w:{edge}"))
            if el is None:
                el = OxmlElement(f"w:{edge}")
                borders.append(el)
            el.set(qn("w:val"), kwargs[edge].get("val", "single"))
            el.set(qn("w:sz"), str(kwargs[edge].get("sz", 6)))
            el.set(qn("w:color"), kwargs[edge].get("color", "000000"))


# ---------------------------------------------------------------- document
doc = Document()

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.0


def page_setup(s, ncols):
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(1.0)
    s.left_margin = Inches(0.625)
    s.right_margin = Inches(0.625)
    set_columns(s, ncols)


page_setup(doc.sections[0], 1)

# ------------------------------------------------ title block (1 column)
p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
p.text = ""
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(p, before=0, after=6)
set_run(p.add_run("Frequency-Query Surrogate Modeling with FiLM "
                  "Conditioning and Passivity-Bounded Output for "
                  "Bandpass Filter Inverse Design"), 20)

for txt, sz, it in [("First A. Author, Second B. Author", 11, False),
                    ("Dept. of Electronic Engineering, Your University, "
                     "City, Country", 10, True),
                    ("{author1, author2}@example.edu", 9, False)]:
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tighten(q, before=0, after=1)
    set_run(q.add_run(txt), sz, italic=it)

tighten(doc.add_paragraph(), after=2)

# ------------------------------------------------ two-column body
sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
page_setup(sec2, 2)

seg_para(doc, [
    ("Abstract—", "bi"),
    ("Surrogate-driven inverse design replaces repeated full-wave "
     "electromagnetic (EM) solver calls with a learned forward model. The "
     "conventional formulation regresses a whole S-parameter curve from a "
     "geometry vector, treating frequency as an output index and offering no "
     "inductive bias toward the oscillatory, resonance-dominated responses it "
     "must reproduce. We present a frequency-query surrogate that instead "
     "learns the pointwise map (", "b"),
    ("x", "bi"), (", ", "b"), ("f", "bi"),
    (") → (S", "b"), ("21", "rs"), (", S", "b"), ("11", "rs"),
    ("), combining a Nyquist-aligned Fourier encoding of the queried "
     "frequency, feature-wise linear modulation generated from the geometry "
     "by a hypernetwork, and a passivity-bounded output head that makes "
     "|S| ≤ 1 structural rather than penalized. We analyse the resulting "
     "benefits—resolution-agnostic querying, a head independent of the "
     "frequency grid, a bias matched to pole–zero physics, and immunity "
     "to the phantom-passband failure mode by which an optimizer adversarially "
     "exploits a surrogate—together with the constrained Top-K inverse "
     "solver built upon it.", "b"),
], size=9)

seg_para(doc, [
    ("Index Terms—", "bi"),
    ("Bandpass filter, feature-wise linear modulation, Fourier features, "
     "frequency query, inverse design, passivity, surrogate modeling.", "b"),
], size=9, before=5)

# ------------------------------------------------ I. Introduction
h1(doc, "I", "Introduction")

body(doc,
     "Full-wave EM simulation is the accuracy standard for planar microwave "
     "filter design, but one swept solution of a multi-resonator layout costs "
     "minutes to hours, and a global search over a nine-dimensional geometry "
     "space at that price is not viable. The established remedy is a neural "
     "surrogate trained on a few hundred solver samples, after which "
     "evolutionary or gradient search runs at millisecond cost [1], [2].",
     indent=0.0)

body(doc,
     "Nearly all such surrogates are whole-curve regressors: a multilayer "
     "perceptron maps the geometry vector onto a fixed-length vector of "
     "S-parameters sampled on a fixed frequency grid. Frequency then exists "
     "only as an output index, with two consequences. First, nothing in the "
     "architecture expresses that adjacent output dimensions are adjacent "
     "frequencies, so the smooth oscillatory structure that resonant responses "
     "possess must be rediscovered from data. Second, the grid is baked into "
     "the output layer, whose width scales with the number of frequency "
     "points; resolution becomes an architectural constant, and resampling the "
     "band requires retraining.")

body(doc,
     "A third consequence is more damaging. An unconstrained regressor can "
     "predict |S21| above 0 dB. Such a value is non-physical for a passive "
     "network, but it is precisely what an inverse optimizer will seek, since "
     "a region where the surrogate promises gain is a region of minimal loss. "
     "The optimizer converges onto a phantom passband that the solver does not "
     "reproduce, and the design loop fails silently. This is no rare numerical "
     "accident: it is the systematic consequence of optimizing against a model "
     "free to be optimistic.")

body(doc,
     "This paper describes a surrogate kernel—a frequency-query network "
     "(FQ-Net)—that removes all three issues by construction, together "
     "with the constrained inverse solver built on it. The frequency-query "
     "idea follows Bi et al. [3]; here it is specialized to a "
     "lumped-parameter rather than pixelated geometry space, and extended with "
     "structural passivity bounding, sharpness-adaptive supervision, and "
     "ensemble-aware inverse search.")

# ------------------------------------------------ II. Problem
h1(doc, "II", "Problem Formulation")
seg_para(doc, [
    ("The device is a planar filter cell: a centre coplanar waveguide (CPW) "
     "carrying an interdigital capacitor, flanked by four rectangular spiral "
     "resonators, with optional cascading of identical cells. Nine geometric "
     "variables are free,", "r"),
], indent=0.0)

equation(doc, [
    [("x", "i"), (" = [", "r"), ("w", "i"), ("0", "rs"), (", ", "r"),
     ("g", "i"), (", ", "r"), ("d", "i"), ("0", "rs"), (", ", "r"),
     ("l", "i"), ("f", "is"), (", ", "r"), ("a", "i"), (", ", "r"),
     ("b", "i"), ("1", "rs"), (", ", "r"), ("b", "i"), ("2", "rs"),
     (", ", "r"), ("d", "i"), (", ", "r"), ("p", "i"), ("0", "rs"),
     ("]", "r"), ("T", "rp"), (" ∈ ℝ", "r"), ("9", "rp")],
], 1)

seg_para(doc, [
    ("namely the CPW conductor width and gap, the centre-gap and finger "
     "lengths of the interdigital capacitor, the spiral arm length, strip "
     "width and strip spacing, and two terms setting the cell period; finger "
     "count, cell count and feed length are held fixed. The response is "
     "sampled at ", "r"), ("N", "i"), ("f", "is"),
    (" = 51 points spanning 0.01–5 GHz and stacked in dB as ", "r"),
    ("y", "i"), (" = [S", "r"), ("21", "rs"), ("; S", "r"), ("11", "rs"),
    ("] ∈ ℝ", "r"), ("2N", "rp"), ("f", "rp"), (".", "r"),
])

body(doc,
     "The feasible set is defined by box bounds together with hard geometric "
     "predicates transcribed from the layout generator, so that every sampled "
     "point is guaranteed to build without self-intersection: a minimum "
     "interdigital strip width, the finger fitting inside the centre gap, a "
     "positive spiral outer centre-line length, board-edge clearance, a "
     "minimum spiral strip spacing, and a 0.1 mm minimum manufacturable "
     "feature on every variable. Training points are drawn by Latin hypercube "
     "sampling with rejection against this set, so the surrogate is never "
     "trained on geometries the process could not realize, and the same "
     "predicates are reused as penalties during inverse search.")

# ------------------------------------------------ III. Surrogate
h1(doc, "III", "The Frequency-Query Surrogate")

h2(doc, "A", "Reformulating the Map")
seg_para(doc, [
    ("FQ-Net learns the pointwise map (", "r"), ("x", "i"), (", ", "r"),
    ("f", "i"), (") → (S", "r"), ("21", "rs"), ("(", "r"), ("f", "i"),
    ("), S", "r"), ("11", "rs"), ("(", "r"), ("f", "i"), ("))", "r"),
    (" rather than ", "r"), ("x", "i"), (" → ℝ", "r"),
    ("2N", "rp"), ("f", "rp"),
    (". A curve is obtained by querying a set of frequencies and "
     "concatenating the replies; in the implementation the grid points are "
     "evaluated as one batched tensor sharing a single set of modulation "
     "coefficients, so a full 102-dimensional curve still costs one forward "
     "pass and the reformulation carries no runtime penalty. Promoting "
     "frequency from an output index to a continuous input is what makes the "
     "three mechanisms below possible.", "r"),
], indent=0.0)

h2(doc, "B", "Nyquist-Aligned Fourier Encoding")
seg_para(doc, [
    ("A raw scalar frequency is a poor network input, as ReLU/GELU networks "
     "are spectrally biased toward low-frequency functions and converge only "
     "slowly on high-frequency detail [4]. The query is normalized to ", "r"),
    ("u", "i"), (" = (", "r"), ("f", "i"), (" − ", "r"), ("f", "i"),
    ("min", "rs"), (")/(", "r"), ("f", "i"), ("max", "rs"), (" − ", "r"),
    ("f", "i"), ("min", "rs"), (") ∈ [0, 1] and lifted by", "r"),
], indent=0.0)

equation(doc, [
    [("γ(", "i"), ("u", "i"), (") = [ sin(", "r"), ("b", "i"),
     ("k", "is"), ("u", "i"), ("), cos(", "r"), ("b", "i"), ("k", "is"),
     ("u", "i"), (") ],   ", "r"), ("k", "i"), (" = 1, …, ", "r"),
     ("K", "i")],
    [("b", "i"), ("k", "is"), (" = π 2", "r"), ("α(k−1)", "rp"),
     (",   α = log", "r"), ("2", "rs"), ("(", "r"), ("N", "i"),
     ("f", "is"), (" /2) / (", "r"), ("K", "i"), (" − 1)", "r")],
], 2)

seg_para(doc, [
    ("The bands are geometrically spaced from ", "r"), ("b", "i"),
    ("1", "rs"), (" = π to ", "r"), ("b", "i"), ("K", "is"),
    (" = π", "r"), ("N", "i"), ("f", "is"),
    (" /2. Tying the fastest basis function to the sampling density of the "
     "training grid keeps it resolvable on that grid and avoids inviting the "
     "network to synthesize oscillations the data cannot support. The choice "
     "is physically motivated: a filter response is a superposition of "
     "resonant poles and transmission zeros, that is, an oscillatory function "
     "of frequency, for which a sinusoidal basis is the natural coordinate "
     "system. ", "r"), ("K", "i"),
    (" trades expressiveness against noise fitting; ", "r"), ("K", "i"),
    (" ∈ [8, 12] is used here, and because the encoding is a fixed "
     "analytic map it adds no trainable parameters.", "r"),
])

h2(doc, "C", "FiLM Conditioning by Hypernetwork")
seg_para(doc, [
    ("The geometry vector is not concatenated with γ(", "r"), ("u", "i"),
    ("). A hypernetwork [5] maps the min–max normalized geometry to a "
     "scale and shift pair for every trunk layer, and these modulate a "
     "frequency-coordinate trunk by feature-wise linear modulation [6]:",
     "r"),
], indent=0.0)

equation(doc, [
    [("(", "r"), ("s", "i"), ("ℓ", "is"), (", ", "r"), ("t", "i"),
     ("ℓ", "is"), (") = Hyper", "r"), ("θ", "rs"), ("(", "r"),
     ("x", "i"), ("),   ℓ = 1, …, ", "r"), ("L", "i")],
    [("h", "i"), ("0", "rs"), (" = φ(", "r"), ("W", "i"),
     ("in", "rs"), (" γ(", "r"), ("u", "i"), ("))", "r")],
    [("h", "i"), ("ℓ", "is"), (" = ", "r"), ("h", "i"),
     ("ℓ−1", "is"), (" + ", "r"), ("D", "i"),
     ("[ φ( (1 + ", "r"), ("s", "i"), ("ℓ", "is"), (") ⊙ ", "r"),
     ("W", "i"), ("ℓ", "is"), (" ", "r"), ("h", "i"),
     ("ℓ−1", "is"), (" + ", "r"), ("t", "i"), ("ℓ", "is"),
     (" ) ]", "r")],
], 3)

seg_para(doc, [
    ("where φ is GELU, ", "r"), ("D", "i"),
    (" is dropout and ⊙ is the elementwise product. The trunk is shared "
     "across all queried frequencies; only the modulation depends on ", "r"),
    ("x", "i"),
    (". Two properties follow. First, conditioning is ", "r"),
    ("multiplicative", "i"),
    (", so a geometry change rescales the entire frequency-feature field "
     "rather than adding a constant offset to it—the appropriate "
     "operation when the physical effect of enlarging a resonator is to "
     "translate and reshape a resonance, not to bias a level. Concatenation, "
     "by contrast, forces the first layer to disentangle two variables of "
     "wholly different character from one flat vector. Second, because the "
     "final hypernetwork layer is zero-initialized, ", "r"),
    ("s", "i"), ("ℓ", "is"), (" = ", "r"), ("t", "i"), ("ℓ", "is"),
    (" = 0 at step zero and (3) reduces to a plain frequency network: "
     "training starts from a parameter-independent mean response and learns "
     "the geometry dependence as a departure from it. This is a stable warm "
     "start that avoids the early-training instability of randomly modulated "
     "trunks, where a large random scale can saturate or extinguish the "
     "signal before any useful gradient is available.", "r"),
])
h2(doc, "D", "Passivity-Bounded Output")
body(doc,
     "The head emits one scalar per output channel and passes it through a "
     "smooth, strictly negative map,", indent=0.0)

equation(doc, [
    [("ŷ(", "i"), ("x", "i"), (", ", "r"), ("f", "i"),
     (") = −softplus(−", "r"), ("W", "i"), ("o", "rs"), (" ", "r"),
     ("h", "i"), ("L", "is"), (") ≤ 0 dB", "r")],
], 4)

seg_para(doc, [
    ("so the prediction is negative in dB for every input whatsoever. A "
     "passive reciprocal two-port satisfies |S| ≤ 1, hence its dB "
     "response cannot exceed zero; (4) makes this a ", "r"),
    ("structural invariant", "i"),
    (" rather than a soft penalty the optimizer can pay. The gradient "
     "σ(−", "r"), ("z", "i"),
    (") is strictly positive everywhere, so unlike a hard clamp—which "
     "zeroes the gradient exactly where the model is most wrong—the "
     "bound never blocks learning. This single line removes the "
     "phantom-passband mode of Section I: no parameter vector, "
     "in-distribution or extrapolated, can be assigned positive gain, so the "
     "inverse optimizer has nothing to exploit.", "r"),
])

h2(doc, "E", "Sharpness-Adaptive Supervision")
body(doc,
     "Training proceeds directly in dB space. Per-bin standardization proved "
     "unusable here: bins whose response is nearly constant across the dataset "
     "have vanishing variance, and dividing by it inflates the normalized "
     "target so that a handful of uninformative bins dominate the loss and the "
     "optimization is effectively hijacked by them. A floor on the per-bin "
     "standard deviation and a −60 dB clip—below which numerical "
     "noise in deep nulls carries no design information yet still inflates "
     "both the per-bin variance and the squared error—are applied "
     "instead.", indent=0.0)

body(doc,
     "A uniform MSE would also spend capacity on the flat stopband, where most "
     "of the output dimensions live, and blur precisely the sharp features on "
     "which a filter is judged. Per-bin weights are therefore derived from the "
     "empirical slope statistics of the training set, and a first-difference "
     "term is added to align predicted and true slopes:")

equation(doc, [
    [("ℒ = ∑", "r"), ("j", "is"), (" ", "r"), ("w", "i"),
     ("j", "is"), (" (ŷ", "r"), ("j", "is"), (" − ", "r"),
     ("y", "i"), ("j", "is"), (")", "r"), ("2", "rp"),
     (" + λ ∑", "r"), ("j", "is"), (" (Δŷ", "r"),
     ("j", "is"), (" − Δ", "r"), ("y", "i"), ("j", "is"), (")", "r"),
     ("2", "rp")],
    [("w", "i"), ("j", "is"), (" ∝ ", "r"), ("c", "i"),
     (" + κ δ", "r"), ("j", "is"), (" / max", "r"), ("i", "rs"),
     (" δ", "r"), ("i", "rs"), (",   ∑", "r"), ("j", "is"),
     (" ", "r"), ("w", "i"), ("j", "is"), (" = 2", "r"), ("N", "i"),
     ("f", "is")],
], 5)

seg_para(doc, [
    ("Here δ", "r"), ("j", "rs"),
    (" is the dataset-mean magnitude of the local first difference at bin ",
     "r"), ("j", "i"),
    (", computed separately over the S", "r"), ("21", "rs"), (" and S", "r"),
    ("11", "rs"),
    (" blocks so that the deeper dynamic range of one does not suppress the "
     "other, and the weights are renormalized to unit mean so the loss scale "
     "is unchanged. The weighting concentrates supervision on transmission "
     "zeros, matching dips and band edges; the difference term is what makes "
     "the predicted skirt slope—the very quantity the inverse objective "
     "later constrains through its roll-off sentinels—trustworthy rather "
     "than merely plausible. An optional Huber form of the first term bounds "
     "the influence of outlier residuals from occasional solver artefacts.",
     "r"),
])

h2(doc, "F", "Training Configuration")
body(doc,
     "The default configuration uses K = 10 Fourier bands and a trunk of "
     "L = 5 residual layers of width H = 256 with dropout 0.05, the "
     "hypernetwork having one hidden layer of the same width. Optimization "
     "uses AdamW with linear warm-up followed by cosine decay, gradient-norm "
     "clipping, a held-out validation split and early stopping on validation "
     "MAE in dB. Five independently seeded members form the ensemble used in "
     "Section IV. Because the entire cost of the method beyond data "
     "acquisition is a few minutes of training, the surrogate is negligible "
     "against the solver sweep that produced the data, which remains the true "
     "bottleneck—and is precisely why sample economy, not training "
     "throughput, is the figure of merit that matters.", indent=0.0)

# ------------------------------------------------ IV. Inverse
h1(doc, "IV", "Constrained Inverse Search")
body(doc,
     "The surrogate is queried inside a derivative-free (CMA-ES [7]) or "
     "gradient-based outer loop. Three elements of that loop are worth stating "
     "because they are what the surrogate design above is meant to serve.",
     indent=0.0)

h2(doc, "A", "Softly Centred, Position-Invariant Passband Objective")
seg_para(doc, [
    ("Specifying a full target curve pins the centre frequency and "
     "over-constrains the search: it demands not merely a compliant filter but "
     "one whose entire response matches a template the designer invented. The "
     "objective instead scores the best admissible passband window. For each "
     "window ", "r"), ("W", "i"),
    ("j", "is"), (" of ", "r"), ("n", "i"), ("w", "is"),
    (" consecutive bins,", "r"),
], indent=0.0)

equation(doc, [
    [("c", "i"), ("j", "is"), (" = mean", "r"), ("i∈Wj", "rs"),
     ("(", "r"), ("P", "i"), (" − ", "r"), ("s", "i"), ("i", "is"),
     (")", "r"), ("+", "rp"), (" + λ", "r"), ("s", "rs"),
     (" mean", "r"), ("i∉Gj", "rs"), ("(", "r"), ("s", "i"), ("i", "is"),
     (" − ", "r"), ("S", "i"), (")", "r"), ("+", "rp")],
    [("+ λ", "r"), ("r", "rs"), (" mean", "r"), ("i∈Rj", "rs"),
     ("(", "r"), ("s", "i"), ("i", "is"), (" − ", "r"), ("R", "i"),
     (")", "r"), ("+", "rp")],
], 6)

seg_para(doc, [
    ("where ", "r"), ("s", "i"), (" is the predicted S", "r"), ("21", "rs"),
    (" in dB, ", "r"), ("P", "i"), (" the passband insertion-loss target, ",
     "r"),
    ("S", "i"), (" the stopband ceiling evaluated outside a guard band ", "r"),
    ("G", "i"), ("j", "is"), (", and ", "r"), ("R", "i"), ("j", "is"),
    (" the two roll-off sentinel bins placed a fixed offset beyond each "
     "passband edge with ceiling ", "r"), ("R", "i"),
    (". The decomposition matters. The first term alone is satisfied by a flat "
     "wideband response, so an optimizer given only that term will happily "
     "return a through-line; the second forbids out-of-band transmission and "
     "the third forces the skirt to fall by a specified amount within a "
     "specified span, and only together do they describe an actual bandpass "
     "shape. A soft centre-frequency constraint admits only those windows "
     "whose centre lies within ", "r"), ("f", "i"), ("0", "rs"),
    (" ± τ, which expresses the specification a designer actually "
     "has—“3.5 GHz, ±0.3 GHz is acceptable”—rather than "
     "an exact template, and leaves the optimizer the freedom to trade centre "
     "frequency for selectivity within the stated tolerance.", "r"),
])

body(doc,
     "The discrete minimum over window position is not differentiable. For "
     "the gradient path it is replaced by a temperature-controlled soft "
     "minimum,")

equation(doc, [
    [("ℒ", "r"), ("pb", "rs"), (" = −τ log ∑", "r"),
     ("j", "is"), (" exp(−", "r"), ("c", "i"), ("j", "is"),
     (" / τ)", "r")],
], 7)

body(doc,
     "so that gradients propagate to the window position as well as to the "
     "response values; inadmissible windows are removed by adding a large "
     "constant to their cost before the aggregation.")

h2(doc, "B", "Robustness to Surrogate Error")
seg_para(doc, [
    ("Three provisions guard the surrogate–solver gap. An ensemble of ",
     "r"), ("M", "i"),
    (" independently seeded networks is averaged, and its inter-member "
     "standard deviation is added to the objective [8], steering the search "
     "away from regions where the surrogate is uncertain—which are "
     "typically the extrapolation regions where it is also wrong. A margin "
     "term tightens ", "r"), ("P", "i"), (" and ", "r"), ("S", "i"),
    (" by a fixed dB offset, so that a solution which is merely marginal under "
     "the surrogate remains compliant under the true response. Geometric and "
     "manufacturability violations enter as a scaled linear penalty. The "
     "complete objective is", "r"),
], indent=0.0)

equation(doc, [
    [("J", "i"), ("(", "r"), ("x", "i"), (") = ℒ", "r"), ("pb", "rs"),
     ("(ŝ(", "r"), ("x", "i"), (")) + λ", "r"), ("u", "rs"),
     (" mean", "r"), ("j", "rs"), (" σ", "r"), ("j", "rs"), ("(", "r"),
     ("x", "i"), (") + ", "r"), ("P", "i"), ("geom", "rs"), ("(", "r"),
     ("x", "i"), (")", "r")],
], 8)

seg_para(doc, [
    ("with ŝ the ensemble-mean S", "r"), ("21", "rs"), (" and σ", "r"),
    ("j", "rs"),
    (" the inter-member standard deviation at bin ", "r"), ("j", "i"),
    (". Note that the uncertainty term is only meaningful because (4) already "
     "removes the systematic optimism: an ensemble of unbounded regressors "
     "can agree confidently on a non-physical prediction, and a variance "
     "penalty does nothing about shared bias.", "r"),
])

h2(doc, "C", "Top-K Diverse Candidates")
body(doc,
     "Because the surrogate optimum is not the true optimum, returning a "
     "single point wastes the one expensive resource available. All restarts "
     "feed a Top-K collector that keeps the K best mutually distinct "
     "candidates: two candidates closer than a threshold in normalized "
     "parameter space are treated as the same solution and merged, keeping the "
     "better, so the returned set spans distinct basins rather than K "
     "near-copies of one. All K are then simulated once in the full-wave "
     "solver and re-ranked by their true loss. K solver calls buy a markedly "
     "higher probability that at least one design is genuinely compliant, at a "
     "cost still negligible against a solver-in-the-loop search, and the "
     "re-ranking makes the final selection independent of surrogate error "
     "entirely.", indent=0.0)

# ------------------------------------------------ V. Properties
h1(doc, "V", "Properties and Benefits")

cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(cap, before=2, after=0)
set_run(cap.add_run("Table I"), 8, small_caps=True)
cap2 = doc.add_paragraph()
cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(cap2, before=0, after=2)
set_run(cap2.add_run("Structural Comparison with a Whole-Curve Regressor"),
        8, small_caps=True)

rows = [
    ("Aspect", "Whole-curve MLP", "FQ-Net"),
    ("Learned map", "x → ℝ²ᴺ", "(x, f) → ℝ²"),
    ("Frequency", "output index", "continuous input"),
    ("Encoding", "implicit", "Fourier, Nyquist-aligned"),
    ("Parameter entry", "input concatenation", "per-layer FiLM"),
    ("Output head", "O(H·Nf) params", "O(H) params"),
    ("Resolution", "fixed at training", "arbitrary at query"),
    ("Passivity", "absent / soft", "structural, ŷ ≤ 0"),
]
tbl = doc.add_table(rows=len(rows), cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
widths = [Inches(0.92), Inches(1.22), Inches(1.28)]
for ri, r in enumerate(rows):
    for ci, txt in enumerate(r):
        c = tbl.cell(ri, ci)
        c.width = widths[ci]
        cell_text(c, txt, size=8, bold=(ri == 0), italic=(ri == 0))
        top = {"val": "single", "sz": 8} if ri == 0 else {"val": "nil"}
        bot = {"val": "single", "sz": 8} if ri in (0, len(rows) - 1) \
            else {"val": "nil"}
        set_cell_border(c, top=top, bottom=bot,
                        left={"val": "nil"}, right={"val": "nil"})
tighten(doc.add_paragraph(), after=1)

h2(doc, "A", "Resolution Decoupled from Architecture")
seg_para(doc, [
    ("The head in (4) has ", "r"), ("O", "i"), ("(", "r"), ("H", "i"),
    (") parameters, against ", "r"), ("O", "i"), ("(", "r"), ("H", "i"),
    ("·", "r"), ("N", "i"), ("f", "is"),
    (") for a whole-curve regressor, and nothing in (2)–(4) depends on ",
     "r"), ("N", "i"), ("f", "is"),
    (" at all. The frequency grid is consequently a query-time argument rather "
     "than an architectural constant: the same trained weights can be "
     "evaluated on a denser grid to resolve a narrow resonance, on a "
     "sub-band, or at a single frequency inside an optimization inner loop, "
     "without retraining and without interpolating a stored curve. Because the "
     "encoding is smooth in ", "r"), ("u", "i"),
    (", the response between training bins is smooth by construction rather "
     "than by post-hoc filtering. This also decouples the data-collection "
     "decision from the modelling decision: a solver sweep taken at one "
     "resolution can be queried at another.", "r"),
], indent=0.0)

h2(doc, "B", "Inductive Bias Matched to the Physics")
body(doc,
     "A bandpass response is characterized by its pole and zero locations. In "
     "a sinusoidal basis, “shift this resonance” is a "
     "low-dimensional change in the modulation coefficients rather than a "
     "coordinated change across many nominally unrelated output units, so the "
     "model expends its capacity on where the features are rather than on "
     "rediscovering that the output is a curve at all. This is precisely where "
     "the present multi-spiral device is hardest for a whole-curve regressor: "
     "it exhibits several in-band matching dips and out-of-band transmission "
     "zeros whose positions move rapidly with geometry, so nearby geometries "
     "produce curves that are similar in shape but displaced in "
     "frequency—the pattern a shift-friendly basis represents cheaply and "
     "an indexed output layer represents expensively.", indent=0.0)

h2(doc, "C", "Structural Physical Consistency")
body(doc,
     "The value of (4) is not that it lowers an average error metric—it "
     "need not—but that it removes an entire failure mode. This "
     "distinction is easy to miss when a surrogate is evaluated only by MAE. "
     "Surrogate-driven optimization does not sample the model uniformly; it "
     "actively searches for the model’s most optimistic prediction, so "
     "any region where the model is optimistic and wrong is exactly the region "
     "the optimizer will find, however small its measure. Bounding the output "
     "at 0 dB eliminates the largest and most systematic such region. The "
     "fraction of predicted S21 samples above 0 dB, a diagnostic worth "
     "reporting for any EM surrogate, is identically zero by construction "
     "rather than small on average.", indent=0.0)

h2(doc, "D", "Training Stability and Sample Economy")
body(doc,
     "Zero-initialized modulation means the model is never worse than a "
     "parameter-independent mean response at initialization, and the training "
     "trajectory monotonically adds geometry dependence rather than having to "
     "recover from a poor random conditioning. Together with dB-space targets, "
     "the variance floor and the −60 dB clip, this removes the loss-scale "
     "pathologies that near-degenerate output bins otherwise cause. "
     "Sharpness-adaptive weighting and slope regularization then direct a "
     "fixed budget of a few hundred solver runs at the features that determine "
     "whether a design meets specification. Since the dataset is the dominant "
     "cost of the whole method, every mechanism that raises accuracy per "
     "sample is worth more than one that raises accuracy per parameter.",
     indent=0.0)

h2(doc, "E", "Drop-In Substitutability")
body(doc,
     "The surrogate exposes the same interface as the whole-curve baseline it "
     "replaces—unit-space forward, ensemble forward, dB prediction and "
     "normalization statistics—and consumes the same dataset and the same "
     "evaluation code, so the forward kernel can be swapped in isolation. This "
     "is a methodological rather than a modelling benefit, and it is the one "
     "that makes the rest assessable: it reduces the comparison between the "
     "two to an ablation of the kernel alone, with data, objective, solver and "
     "metric held fixed, which is the only basis on which a difference in "
     "end-to-end design quality can be attributed to the model rather than to "
     "the experimental setup.", indent=0.0)

# ------------------------------------------------ VI. Conclusion
h1(doc, "VI", "Conclusion")
body(doc,
     "We described a frequency-query surrogate for planar bandpass filter "
     "inverse design in which frequency is a continuous, Fourier-encoded "
     "input, geometry enters as hypernetwork-generated FiLM modulation of a "
     "shared frequency trunk, and passivity is enforced structurally by a "
     "bounded output head. The formulation makes frequency resolution a "
     "query-time choice, aligns the model’s inductive bias with the "
     "oscillatory physics of resonant filters, and eliminates the non-physical "
     "predictions that a surrogate-driven optimizer would otherwise exploit. "
     "Combined with a softly centred passband objective carrying explicit "
     "roll-off sentinels, ensemble-uncertainty regularization, manufacturing "
     "constraints and Top-K re-ranking against the full-wave solver, it forms "
     "a design loop whose principal failure modes are suppressed structurally "
     "rather than statistically. Quantitative comparison against a whole-curve "
     "baseline on a common dataset, using the shared evaluation path of "
     "Section V-E, is the subject of ongoing work.", indent=0.0)

# ------------------------------------------------ References
h1(doc, "", "References")
refs = [
    "Q.-J. Zhang, K. C. Gupta, and V. K. Devabhaktuni, “Artificial neural "
    "networks for RF and microwave design—from theory to practice,” "
    "IEEE Trans. Microw. Theory Techn., vol. 51, no. 4, pp. 1339–1350, "
    "Apr. 2003.",
    "J. E. Rayas-Sánchez, “EM-based optimization of microwave circuits "
    "using artificial neural networks: The state-of-the-art,” IEEE "
    "Trans. Microw. Theory Techn., vol. 52, no. 1, pp. 420–435, "
    "Jan. 2004.",
    "J. Bi, X. Zhou, J. Xia, S. Chen, and W. S. Chan, “Frequency-query "
    "enhanced electromagnetic surrogate modeling with edge anti-aliasing "
    "pixelation for bandpass filter inverse design,” in Proc. IEEE MTT-S "
    "Int. Microw. Symp. (IMS), Jun. 2025, pp. 614–617.",
    "M. Tancik et al., “Fourier features let networks learn high "
    "frequency functions in low dimensional domains,” in Adv. Neural Inf. "
    "Process. Syst. (NeurIPS), vol. 33, 2020, pp. 7537–7547.",
    "D. Ha, A. Dai, and Q. V. Le, “HyperNetworks,” in Proc. Int. "
    "Conf. Learn. Represent. (ICLR), 2017.",
    "E. Perez, F. Strub, H. de Vries, V. Dumoulin, and A. Courville, "
    "“FiLM: Visual reasoning with a general conditioning layer,” in "
    "Proc. AAAI Conf. Artif. Intell., 2018, pp. 3942–3951.",
    "N. Hansen, “The CMA evolution strategy: A tutorial,” "
    "arXiv:1604.00772, 2016.",
    "B. Lakshminarayanan, A. Pritzel, and C. Blundell, “Simple and "
    "scalable predictive uncertainty estimation using deep ensembles,” in "
    "Adv. Neural Inf. Process. Syst. (NeurIPS), vol. 30, 2017, "
    "pp. 6402–6413.",
]
for i, r in enumerate(refs, 1):
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    tighten(q, before=0, after=0)
    q.paragraph_format.left_indent = Inches(0.20)
    q.paragraph_format.first_line_indent = Inches(-0.20)
    set_run(q.add_run(f"[{i}]\t{r}"), 8)

doc.save(OUT)
print(f"saved -> {OUT}")
